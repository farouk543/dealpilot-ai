import io

from openpyxl import Workbook


def build_xlsx(deal: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Résumé"
    prop = deal.get("property") or {}
    ws.append(["Dossier", deal.get("deal_id", "")])
    ws.append(["Type de bien", prop.get("property_type")])
    ws.append(["Localisation", prop.get("location")])
    ws.append(["Adresse", prop.get("address")])
    ws.append(["Étape actuelle", deal.get("current_step")])

    underwriting = deal.get("underwriting")
    if underwriting and underwriting.get("scenarios"):
        ws2 = wb.create_sheet("Souscription")
        ws2.append(["Scénario", "NOI", "Cap rate %", "Cash-flow", "Cash-on-cash %", "DSCR", "IRR %"])
        for s in underwriting["scenarios"]:
            ws2.append(
                [
                    s.get("name"),
                    s.get("noi"),
                    s.get("cap_rate_pct"),
                    s.get("cash_flow_before_tax"),
                    s.get("cash_on_cash_pct"),
                    s.get("dscr") if isinstance(s.get("dscr"), (int, float)) else str(s.get("dscr")),
                    s.get("irr_pct"),
                ]
            )

    risk_register = deal.get("risk_register")
    if risk_register and risk_register.get("risks"):
        ws3 = wb.create_sheet("Risques")
        ws3.append(["Sévérité", "Catégorie", "Titre", "Preuve", "Impact", "Action suivante"])
        for r in risk_register["risks"]:
            ws3.append(
                [r.get("severity"), r.get("category"), r.get("title"), r.get("evidence"), r.get("impact"), r.get("next_action")]
            )

    dd = deal.get("due_diligence")
    if dd and dd.get("tasks"):
        ws4 = wb.create_sheet("Due diligence")
        ws4.append(["Type", "Description", "Sévérité", "Statut", "Risque lié"])
        for t in dd["tasks"]:
            ws4.append(
                [t.get("task_type"), t.get("description"), t.get("severity_driver"), t.get("status"), t.get("linked_risk_title")]
            )

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
